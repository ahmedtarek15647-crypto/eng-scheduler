"""
Export utilities for Engineering Project Scheduler.
Generates Excel spreadsheets (.xlsx), WhatsApp proposal messages,
and formatted project briefs with full support for project codes, completion dates, and cash tracking.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_schedule(project_data):
    """
    Builds a professional, multi-sheet Excel file (.xlsx) with:
    - Sheet 1: Project Executive Summary, Cash Status & Milestones
    - Sheet 2: Detailed Engineering Work Breakdown Schedule (WBS)
    """
    wb = Workbook()
    
    currency = project_data.get("currency") or "EGP"
    is_fin = bool(project_data.get("is_finished", 0))
    is_paid = bool(project_data.get("is_cash_received", 0))
    
    actual_end = project_data.get("actual_end_date") or ("Completed" if is_fin else "In Progress")
    cash_date = project_data.get("received_cash_date") or ("Received" if is_paid else "Pending")

    # ---------------- SHEET 1: SUMMARY & MILESTONES ----------------
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    # Styling Palettes
    navy_header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    soft_gray_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    title_font = Font(name="Calibri", size=16, bold=True, color="1B365D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    status_paid_font = Font(name="Calibri", size=11, bold=True, color="065F46")
    status_pending_font = Font(name="Calibri", size=11, bold=True, color="92400E")

    thin_border = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )

    code = project_data.get("project_code", "N/A")
    # Title Banner
    ws1["B2"] = f"PROJECT REPORT: {code} - {project_data.get('name', 'N/A')}"
    ws1["B2"].font = title_font
    ws1["B3"] = f"Client: {project_data.get('client_name', 'Client')} | Generated via EngScheduler Pro"
    ws1["B3"].font = Font(name="Calibri", size=11, italic=True, color="475467")

    # Status Strings
    fin_status_str = f"FINISHED (Completed: {actual_end})" if is_fin else "IN PROGRESS"
    cash_status_str = f"PAID IN FULL (Received: {cash_date})" if is_paid else "PAYMENT PENDING"

    # Metadata Grid
    metadata = [
        ("Project Code:", code),
        ("Project Title:", project_data.get("name", "N/A")),
        ("Client Name:", project_data.get("client_name", "N/A")),
        ("Client Contact:", project_data.get("client_contact", "N/A")),
        ("Discipline:", project_data.get("discipline", "N/A")),
        ("Start Date:", project_data.get("start_date", "N/A")),
        ("Scheduled End Date:", project_data.get("scheduled_end_date", "N/A")),
        ("Execution Status:", fin_status_str),
        ("Project Cash (100%):", f"{project_data.get('cost', 0):,.2f} {currency}"),
        ("Cash Collection Status:", cash_status_str),
        ("Weekend Config:", project_data.get("weekend_type", "Fri-Sat")),
        ("Project Remarks:", project_data.get("remarks", "None"))
    ]

    for idx, (label, val) in enumerate(metadata, start=5):
        ws1[f"B{idx}"] = label
        ws1[f"B{idx}"].font = bold_font
        ws1[f"B{idx}"].fill = soft_gray_fill
        ws1[f"B{idx}"].border = thin_border
        
        ws1[f"C{idx}"] = str(val)
        ws1[f"C{idx}"].font = normal_font
        ws1[f"C{idx}"].border = thin_border

        # Highlight status rows
        if label == "Execution Status:":
            ws1[f"C{idx}"].fill = green_fill if is_fin else amber_fill
            ws1[f"C{idx}"].font = status_paid_font if is_fin else status_pending_font
        elif label == "Cash Collection Status:":
            ws1[f"C{idx}"].fill = green_fill if is_paid else amber_fill
            ws1[f"C{idx}"].font = status_paid_font if is_paid else status_pending_font

    # Milestone Invoicing Section
    m_row_start = 18
    ws1[f"B{m_row_start}"] = "PAYMENT & INVOICING MILESTONES"
    ws1[f"B{m_row_start}"].font = Font(name="Calibri", size=13, bold=True, color="1B365D")

    m_headers = ["Milestone Title", "Trigger Deliverable", "%", f"Amount ({currency})", "Target / Due Date", "Payment Status"]
    cols = ["B", "C", "D", "E", "F", "G"]

    for col_letter, header_text in zip(cols, m_headers):
        cell = ws1[f"{col_letter}{m_row_start + 1}"]
        cell.value = header_text
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    curr_row = m_row_start + 2
    for m in project_data.get("milestones", []):
        # If project is marked as cash received, all milestones show Paid!
        if is_paid:
            m_status = "Paid"
            m_due = cash_date
        else:
            m_status = m.get("status", "Pending")
            m_due = m.get("due_date", "")

        ws1[f"B{curr_row}"] = m.get("milestone", "")
        ws1[f"C{curr_row}"] = m.get("trigger", "")
        ws1[f"D{curr_row}"] = f"{m.get('percentage', 0)}%"
        ws1[f"E{curr_row}"] = f"{m.get('amount', 0):,.2f} {currency}"
        ws1[f"F{curr_row}"] = m_due
        ws1[f"G{curr_row}"] = m_status

        for c in cols:
            ws1[f"{c}{curr_row}"].font = normal_font
            ws1[f"{c}{curr_row}"].border = thin_border
            
        ws1[f"D{curr_row}"].alignment = Alignment(horizontal="center")
        ws1[f"E{curr_row}"].alignment = Alignment(horizontal="right")
        ws1[f"F{curr_row}"].alignment = Alignment(horizontal="center")
        ws1[f"G{curr_row}"].alignment = Alignment(horizontal="center")
        
        # Color status cell
        if m_status == "Paid":
            ws1[f"G{curr_row}"].fill = green_fill
            ws1[f"G{curr_row}"].font = status_paid_font
        else:
            ws1[f"G{curr_row}"].fill = amber_fill
            ws1[f"G{curr_row}"].font = status_pending_font

        curr_row += 1

    # ---------------- SHEET 2: WBS SCHEDULE ----------------
    ws2 = wb.create_sheet(title="Engineering WBS Schedule")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"WORK BREAKDOWN STRUCTURE (WBS) - {code} {project_data.get('name', 'Project')}"
    ws2["A1"].font = title_font

    wbs_headers = ["ID", "Phase", "Task / Activity", "Start Date", "End Date", "Working Days", "Key Deliverable", f"Budget ({currency})", "Execution Status"]
    wbs_cols = ["A", "B", "C", "D", "E", "F", "G", "H", "I"]

    for col_letter, h in zip(wbs_cols, wbs_headers):
        cell = ws2[f"{col_letter}3"]
        cell.value = h
        cell.font = header_font
        cell.fill = navy_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    r_idx = 4
    for t in project_data.get("tasks", []):
        # If project is marked finished, all tasks reflect Completed!
        if is_fin:
            t_status = "Completed"
        else:
            t_status = t.get("status", "Not Started")

        ws2[f"A{r_idx}"] = t.get("id", "")
        ws2[f"B{r_idx}"] = t.get("phase", "")
        ws2[f"C{r_idx}"] = t.get("name", "")
        ws2[f"D{r_idx}"] = t.get("start", "")
        ws2[f"E{r_idx}"] = t.get("end", "")
        ws2[f"F{r_idx}"] = t.get("working_days", "")
        ws2[f"G{r_idx}"] = t.get("deliverable", "")
        ws2[f"H{r_idx}"] = f"{t.get('budget', 0):,.2f} {currency}"
        ws2[f"I{r_idx}"] = t_status

        for c in wbs_cols:
            ws2[f"{c}{r_idx}"].font = normal_font
            ws2[f"{c}{r_idx}"].border = thin_border
            
        ws2[f"A{r_idx}"].alignment = Alignment(horizontal="center")
        ws2[f"D{r_idx}"].alignment = Alignment(horizontal="center")
        ws2[f"E{r_idx}"].alignment = Alignment(horizontal="center")
        ws2[f"F{r_idx}"].alignment = Alignment(horizontal="center")
        ws2[f"H{r_idx}"].alignment = Alignment(horizontal="right")
        ws2[f"I{r_idx}"].alignment = Alignment(horizontal="center")
        
        if t_status == "Completed":
            ws2[f"I{r_idx}"].fill = green_fill
            ws2[f"I{r_idx}"].font = status_paid_font
        elif t_status == "In Progress":
            ws2[f"I{r_idx}"].fill = amber_fill
            ws2[f"I{r_idx}"].font = status_pending_font

        r_idx += 1

    # Auto-adjust column widths for all sheets
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def format_whatsapp_proposal(project_data):
    """
    Formats a comprehensive mobile proposal and status message with correct currency.
    """
    code = project_data.get('project_code', '')
    currency = project_data.get('currency') or "EGP"
    cost = project_data.get('cost', 0)
    scheduled_end = project_data.get('scheduled_end_date') or project_data.get('deadline', 'N/A')
    
    is_fin = bool(project_data.get("is_finished", 0))
    is_cash = bool(project_data.get("is_cash_received", 0))

    fin_text = f"✅ Finished ({project_data.get('actual_end_date', 'Completed')})" if is_fin else "⏳ In Progress"
    cash_text = f"💰 Cash Received ({project_data.get('received_cash_date', 'Paid')})" if is_cash else "⏳ Payment Pending"

    lines = [
        f"🏛️ *PROJECT BRIEF & SCHEDULE*",
        f"━━━━━━━━━━━━━━━━━━━━━━",
        f"🔖 *Project Code:* {code}",
        f"📋 *Project:* {project_data.get('name', 'N/A')}",
        f"👤 *Client:* {project_data.get('client_name', 'N/A')}",
        f"⚙️ *Discipline:* {project_data.get('discipline', 'N/A')}",
        f"🗓️ *Start Date:* {project_data.get('start_date', 'N/A')}",
        f"🏁 *Scheduled End:* {scheduled_end}",
        f"💵 *Project Cash (100%):* {cost:,.2f} {currency}",
        f"📌 *Status:* {fin_text} | {cash_text}",
        f"",
        f"📌 *KEY DELIVERABLES:*",
    ]

    for t in project_data.get("tasks", []):
        t_status = " [✅ Completed]" if is_fin else ""
        lines.append(f"• *{t.get('phase', '')}* (By: {t.get('end')}){t_status}")
        lines.append(f"   ↳ {t.get('deliverable')}")

    lines.append("")
    lines.append(f"💳 *PAYMENT MILESTONES:*")
    for m in project_data.get("milestones", []):
        m_status = " [✅ Paid]" if is_cash else f" (Due: {m.get('due_date')})"
        lines.append(f"• {m.get('milestone')}: *{m.get('percentage')}%* ({m.get('amount', 0):,.2f} {currency}){m_status}")

    remarks = project_data.get('remarks', '').strip()
    if remarks:
        lines.append("")
        lines.append(f"📝 *Remarks / Notes:*")
        lines.append(f"{remarks}")

    lines.append("")
    lines.append("━━━━━━━━━━━━━━━━━━━━━━")
    lines.append("Thank you for your business! Please let us know if any questions arise.")
    
    return "\n".join(lines)
